import os 
from llm_analysis import generate_placeholders
from docx import Document
import cloudconvert
import requests
import convertapi
import glob
from docx.shared import Pt, RGBColor


OUTPUT_PATH = os.environ["OUTPUT_PATH"]
os.makedirs(OUTPUT_PATH, exist_ok=True)

cloudconvert.configure(api_key=os.environ["CLOUDCONVERT_API_KEY"])

CV_TEMPLATE_PATH = os.environ.get("CV_TEMPLATE_PATH", "cv_template.docx")

def __replace_in_doc(doc, placeholders: dict):
    #print("DEBUG: print template")
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            #print(run.text)
            for key, value in placeholders.items():
                if f"{{{{{key}}}}}" in run.text:
                    run.text = run.text.replace(f"{{{{{key}}}}}", value)

                    if key == "ROLE":
                        run.bold = True
                        run.font.size = Pt(14)
                        run.font.color.rgb = RGBColor(0x1B, 0x73, 0xC8)


def __convert_to_pdf(doc_path, pdf_path):
    print("Converting to pdf...")
    try:
        # CloudConvert
        job = cloudconvert.Job.create(payload={
            "tasks": {
                "upload": {"operation": "import/upload"},
                "convert": {
                    "operation": "convert",
                    "input": "upload",
                    "input_format": "docx",
                    "output_format": "pdf"
                },
                "export": {
                    "operation": "export/url",
                    "input": "convert"
                }
            }
        })
        upload_task = next(t for t in job["tasks"] if t["name"] == "upload")
        cloudconvert.Task.upload(file_name=doc_path, task=upload_task)
        job = cloudconvert.Job.wait(id=job["id"])
        export_task = next(t for t in job["tasks"] if t["name"] == "export")
        url = export_task["result"]["files"][0]["url"]
        response = requests.get(url)
        with open(pdf_path, "wb") as f:
            f.write(response.content)

    except Exception as e:
        print(f"CloudConvert failed ({e}), falling back to ConvertAPI...")
        convertapi.api_credentials = os.environ["CONVERTAPI_SECRET"]
        convertapi.convert("pdf", {"File": doc_path}, from_format="docx").save_files(
            os.path.dirname(pdf_path)
        )

        saved = glob.glob(os.path.join(os.path.dirname(pdf_path), "*.pdf"))[0]
        os.rename(saved, pdf_path)


def customize_cv(offer, template_path, output_path, name, candidate_profile, competencies, libraries, languages, tools)->str:

    # Generate placeholders using candidate pool and job description
    placeholders = generate_placeholders(offer, candidate_profile=candidate_profile, competencies=competencies, libraries=libraries, languages=languages, tools=tools)

    # open template docx
    doc = Document(template_path)

    # replace 
    __replace_in_doc(doc=doc, placeholders=placeholders)

    # save docx
    company = offer["company"].replace(" ", "_").lower()
    filename = "CV_" + name.replace(" ", "_").lower() + "_" +company + ".docx"
    doc_path = output_path + filename
    doc.save(doc_path)

    # convert and save pdf
    pdf_path = doc_path.replace(".docx", ".pdf")

    __convert_to_pdf(doc_path, pdf_path)

    return pdf_path